"""Supporting outputs: slim results database, Excel workbook, figures, dashboard JSON."""
from __future__ import annotations

import json
import sqlite3
from pathlib import Path

import numpy as np
import pandas as pd

from . import db
from .viz import COLORS, NAMES, ORDER, SEGMENT_COLORS

AUTHOR = "Sasmitha Jayakody"          # workbook and PDF metadata

ROOT = db.ROOT
OUT = ROOT / "outputs"
FIG = ROOT / "docs" / "figures"


def read_results(db_path=None) -> dict[str, pd.DataFrame]:
    con = sqlite3.connect(db_path or db.DEFAULT_DB)
    names = [r[0] for r in con.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND (name LIKE 'res_%' OR name='securities')")]
    out = {n.replace("res_", ""): pd.read_sql(f"SELECT * FROM {n}", con) for n in names}
    con.close()
    return out


def export_results_db(db_path=None, dest: Path | None = None) -> Path:
    """Copy res_* tables and the security reference into a small file for the dashboard."""
    dest = dest or (ROOT / "data" / "results.db")
    if dest.exists():
        dest.unlink()
    src = sqlite3.connect(db_path or db.DEFAULT_DB)
    dst = sqlite3.connect(dest)
    names = [r[0] for r in src.execute(
        "SELECT name FROM sqlite_master WHERE type='table' AND (name LIKE 'res_%' OR name='securities')")]
    for n in names:
        pd.read_sql(f"SELECT * FROM {n}", src).to_sql(n, dst, index=False)
    dst.commit()
    dst.close()
    src.close()
    return dest


# --------------------------------------------------------------------------- excel
def build_excel(R: dict[str, pd.DataFrame], path: Path | None = None) -> Path:
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    path = path or (OUT / "etf_portfolio_results.xlsx")
    path.parent.mkdir(parents=True, exist_ok=True)
    meta = dict(zip(R["meta"]["key"], R["meta"]["value"]))
    S = R["summary"].set_index("code").reindex(ORDER)

    sheets: dict[str, pd.DataFrame] = {}
    sheets["Summary"] = (
        S[["name", "rule", "cagr", "vol", "sharpe", "sortino", "max_dd", "dd_trough", "dd_recovery", "calmar",
           "tracking_error", "info_ratio", "beta", "up_capture", "down_capture", "var95_daily", "cvar95_daily",
           "turnover_ann", "cost_bps_ann", "gross_cagr", "pct_pos_months"]]
        .rename(columns={
            "name": "Portfolio", "rule": "Rebalancing rule", "cagr": "Return p.a.", "vol": "Volatility p.a.",
            "sharpe": "Sharpe", "sortino": "Sortino", "max_dd": "Max drawdown", "dd_trough": "Trough",
            "dd_recovery": "Recovered", "calmar": "Calmar", "tracking_error": "Tracking error",
            "info_ratio": "Information ratio", "beta": "Beta", "up_capture": "Up capture",
            "down_capture": "Down capture", "var95_daily": "VaR 95% (daily)", "cvar95_daily": "CVaR 95% (daily)",
            "turnover_ann": "Turnover p.a.", "cost_bps_ann": "Costs p.a. (bp)", "gross_cagr": "Gross return p.a.",
            "pct_pos_months": "% positive months"}).T
    )
    sheets["Calendar returns"] = R["calendar"].pivot(index="year", columns="code", values="return")[ORDER]
    sheets["Regime returns"] = R["periods"].pivot(index="period", columns="code", values="cagr")[ORDER]
    for dim in ("sector", "country", "currency", "segment"):
        t = R["exposures"][R["exposures"].dimension == dim].pivot_table(
            index="bucket", columns="code", values="weight", fill_value=0.0)
        sheets[f"Exposure {dim}"] = t.reindex(columns=ORDER).sort_values("SAA", ascending=False)
    sheets["Duration"] = R["duration"].set_index("code").reindex(ORDER).T
    sheets["Attribution totals"] = R["brinson_totals"].set_index(["code", "period"])
    sheets["Attribution segments"] = R["brinson_segments"][R["brinson_segments"].period == "Full period"] \
        .set_index(["code", "segment"]).drop(columns="period")
    sheets["Factor model"] = R["factor"].set_index(["code", "mode", "factor"])
    sheets["Risk contributions"] = R["risk_contrib"].pivot_table(
        index=["code", "ticker"], columns="window", values=["weight", "pct_risk"])
    sheets["Stress historical"] = R["stress_replay"].pivot_table(
        index="window", columns="code", values=["replay_return", "realised_return"])
    sheets["Stress scenarios"] = R["stress_scenarios"].pivot(index="scenario", columns="code", values="return")[ORDER]
    sheets["Rebalancing rules"] = R["rebalance_study"].set_index("rule")
    sheets["Cost sensitivity"] = R["cost_sensitivity"].pivot(index="cost_multiplier", columns="code",
                                                            values=["cagr", "sharpe", "cost_bps_ann"])
    w = R["weights"]
    sheets["Current weights"] = w[w.date == w.date.max()].pivot_table(index="ticker", columns="code", values="weight")
    sheets["Notes"] = pd.DataFrame({
        "item": ["Source", "Backtest window", "Prices as of", "Pipeline run", "Holdings as of",
                 "Transaction cost model", "Risk-free rate", "Caveat"],
        "value": [meta["data"], f"{meta['backtest_start']} to {meta['backtest_end']}", meta["prices_end"],
                  meta["run_at"], meta.get("holdings_as_of", ""),
                  "one-way bps per ETF from data/reference/etf_universe.csv, charged on traded weight",
                  "SHV (0-1y Treasury ETF) total return",
                  "Values are outputs of a backtest, not live performance. See docs/methodology.md."]
    }).set_index("item")

    with pd.ExcelWriter(path, engine="openpyxl") as xl:
        for name, df in sheets.items():
            df.to_excel(xl, sheet_name=name[:31])
    # formatting
    from openpyxl import load_workbook
    wb = load_workbook(path)
    head = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    body = Font(name="Arial", size=10)
    fill = PatternFill("solid", fgColor="1F3864")
    thin = Border(bottom=Side(style="thin", color="D9D9D9"))
    for ws in wb.worksheets:
        ws.freeze_panes = "B2"
        for cell in ws[1]:
            cell.font = head
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body
                cell.border = thin
                if isinstance(cell.value, float):
                    v = abs(cell.value)
                    cell.number_format = "0.00" if v > 1.5 else "0.0%"
        widths = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    widths[cell.column] = max(widths.get(cell.column, 10), min(len(str(cell.value)) + 2, 42))
        for col, width in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.row_dimensions[1].height = 30
    wb.properties.creator = AUTHOR
    wb.properties.lastModifiedBy = AUTHOR
    wb.save(path)
    return path


# --------------------------------------------------------------------------- figures
def _mpl():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "DejaVu Sans", "font.size": 8, "axes.edgecolor": "#c3c2b7", "axes.labelcolor": "#52514e",
        "axes.titlesize": 9, "axes.titleweight": "bold", "axes.titlecolor": "#0b0b0b", "axes.grid": True,
        "grid.color": "#e1e0d9", "grid.linewidth": 0.6, "axes.spines.top": False, "axes.spines.right": False,
        "xtick.color": "#898781", "ytick.color": "#898781", "figure.facecolor": "white", "axes.facecolor": "white",
        "legend.frameon": False, "legend.fontsize": 7.5,
    })
    return plt


def build_figures(R: dict[str, pd.DataFrame]) -> list[Path]:
    plt = _mpl()
    FIG.mkdir(parents=True, exist_ok=True)
    paths = []
    nav = R["nav"].copy()
    nav["date"] = pd.to_datetime(nav["date"])

    fig, (ax1, ax2) = plt.subplots(2, 1, figsize=(6.6, 4.0), height_ratios=[2, 1], sharex=True)
    for c in ORDER:
        x = nav[nav.code == c]
        ax1.plot(x.date, 100 * x.nav, color=COLORS[c], lw=1.6, label=NAMES[c])
        ax2.plot(x.date, 100 * x.drawdown, color=COLORS[c], lw=1.1)
    ax1.set_title("Growth of $100, net of costs")
    ax1.legend(loc="upper left", ncols=2)
    ax2.set_title("Drawdown (%)")
    fig.tight_layout()
    p = FIG / "performance.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    paths.append(p)

    tot = R["brinson_totals"]
    full = tot[tot.period == "Full period"].set_index("code")
    fig, ax = plt.subplots(figsize=(6.6, 2.4))
    labels = ["Allocation", "Selection", "Costs", "Benchmark residual", "Total active"]
    width = 0.26
    for i, c in enumerate(["SAA", "MINVAR", "QV"]):
        r = full.loc[c]
        vals = [r.allocation, r.selection, r.costs, r.bench_residual, r.active]
        ax.bar(np.arange(5) + (i - 1) * width, [100 * v for v in vals], width * 0.9, color=COLORS[c], label=NAMES[c])
    ax.set_xticks(range(5), labels)
    ax.set_ylabel("cumulative, %")
    ax.set_title("Active return versus benchmark, decomposed")
    ax.legend(ncols=3)
    fig.tight_layout()
    p = FIG / "attribution.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    paths.append(p)

    rp = R["stress_replay"]
    piv = rp.pivot(index="window", columns="code", values="replay_return").reindex(columns=ORDER)
    order = list(rp.window.unique())
    piv = piv.reindex(order)
    fig, ax = plt.subplots(figsize=(6.6, 3.2))
    y = np.arange(len(piv))
    for i, c in enumerate(ORDER):
        ax.barh(y + (1.5 - i) * 0.2, 100 * piv[c], 0.18, color=COLORS[c], label=NAMES[c])
    ax.set_yticks(y, piv.index, fontsize=7)
    ax.invert_yaxis()
    ax.set_xlabel("return over episode, %")
    ax.set_title("Historical stress replay, current weights")
    ax.legend(ncols=2)
    fig.tight_layout()
    p = FIG / "stress.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    paths.append(p)

    rc = R["risk_contrib"]
    rc = rc[rc.window == "1y"].groupby(["code", "segment"], as_index=False)[["weight", "pct_risk"]].sum()
    fig, ax = plt.subplots(figsize=(6.6, 3.1))
    ypos, labels, seen = [], [], {}
    for i, c in enumerate(ORDER):
        for j, (lbl, col) in enumerate((("capital", "weight"), ("risk", "pct_risk"))):
            left = 0.0
            yy = i * 2.4 + j
            for seg in SEGMENT_COLORS:
                v = rc[(rc.code == c) & (rc.segment == seg)][col]
                if v.empty or abs(float(v.iloc[0])) < 1e-6:
                    continue
                bar = ax.barh(yy, 100 * float(v.iloc[0]), 0.8, left=left, color=SEGMENT_COLORS[seg],
                              edgecolor="white", linewidth=0.8)
                seen.setdefault(seg, bar)
                left += 100 * float(v.iloc[0])
            ypos.append(yy)
            labels.append(f"{NAMES[c]} · {lbl}")
    ax.set_yticks(ypos, labels, fontsize=7)
    ax.invert_yaxis()
    ax.set_xlabel("% of portfolio / % of volatility")
    ax.set_title("Capital versus risk")
    ax.legend(seen.values(), seen.keys(), ncols=5, fontsize=6.5, loc="upper center",
              bbox_to_anchor=(0.5, -0.22))
    fig.tight_layout()
    p = FIG / "risk_contribution.png"
    fig.savefig(p, dpi=200)
    plt.close(fig)
    paths.append(p)
    return paths


# --------------------------------------------------------------------------- dashboard json
def _clean(df: pd.DataFrame) -> pd.DataFrame:
    """NaN is not valid JSON; the dashboard treats null as 'not available'."""
    return df.astype(object).where(pd.notna(df), None)


def _rolling_payload(roll: pd.DataFrame) -> dict:
    out: dict[str, dict] = {}
    for metric, g in roll.groupby("metric"):
        out[metric] = {}
        for code, gg in g.groupby("code"):
            out[metric][code] = {
                "dates": gg["date"].tolist(),
                "value": [None if pd.isna(v) else round(float(v), 4) for v in gg["value"]],
            }
    return out


def dashboard_json(R: dict[str, pd.DataFrame], path: Path | None = None) -> Path:
    """Compact payload for the static web dashboard."""
    path = path or (OUT / "dashboard_data.json")
    R = {k: (_clean(v) if isinstance(v, pd.DataFrame) else v) for k, v in R.items()}
    nav = R["nav"].copy()
    nav["date"] = pd.to_datetime(nav["date"])
    nav = nav[nav.date.dt.weekday == 4]        # weekly points keep the file small
    payload = {
        "meta": dict(zip(R["meta"]["key"], R["meta"]["value"])),
        "summary": R["summary"].to_dict("records"),
        "nav": {c: {"dates": nav[nav.code == c].date.dt.strftime("%Y-%m-%d").tolist(),
                    "nav": [round(v, 4) for v in nav[nav.code == c].nav],
                    "dd": [round(v, 4) for v in nav[nav.code == c].drawdown]} for c in ORDER},
        "calendar": R["calendar"].to_dict("records"),
        "periods": R["periods"].to_dict("records"),
        "exposures": R["exposures"].to_dict("records"),
        "duration": R["duration"].to_dict("records"),
        "top_holdings": R["top_holdings"].to_dict("records"),
        "brinson_totals": R["brinson_totals"].to_dict("records"),
        "brinson_segments": R["brinson_segments"][R["brinson_segments"].period == "Full period"].to_dict("records"),
        "factor": R["factor"].to_dict("records"),
        "factor_meta": R["factor_meta"].to_dict("records"),
        "risk_contrib": R["risk_contrib"][R["risk_contrib"].window == "1y"]
            .groupby(["code", "segment"], as_index=False)[["weight", "pct_risk"]].sum().to_dict("records"),
        "stress_replay": R["stress_replay"].to_dict("records"),
        "stress_scenarios": R["stress_scenarios"].to_dict("records"),
        "rebalance_study": R["rebalance_study"].to_dict("records"),
        "cost_sensitivity": R["cost_sensitivity"].to_dict("records"),
        "turnover_year": R["turnover_year"].to_dict("records"),
        "sensitivities": R["sensitivities"].to_dict("records"),
        "current_weights": R["weights"][R["weights"].date == R["weights"].date.max()].to_dict("records"),
        "rolling": _rolling_payload(R["rolling"]),
        "stock_bond_corr": {"dates": R["stock_bond_corr"].date.tolist(),
                            "value": [None if pd.isna(v) else round(v, 3) for v in R["stock_bond_corr"].value]},
    }
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, separators=(",", ":"), allow_nan=False), encoding="utf-8")
    return path
